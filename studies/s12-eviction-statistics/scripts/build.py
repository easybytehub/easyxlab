"""S12 step 1: from the locked files to the tables the analysis uses.

Writes (aggregates only):
  data/districts.csv          431 districts: province, phase, weights, practised evictions 2023-2025, flags
  data/province_exposure.csv  50 provinces: baseline B_p and phase shares W1, W2, W3
  data/province_quarter.csv   province x quarter x series, 2021Q1-2026Q1 (CGPJ series file, latest vintage)
  data/cgpj_notes.csv         the CGPJ's own missing/estimated-data notes, literal, matched to districts
  work/panel.json             full series 2007/2013-2026Q1 by province and TSJ, D1 by district (not published)
"""
from __future__ import annotations

import re
from collections import defaultdict

from s12lib import (CEUTA_MELILLA_HOST, DATA, PHASE_QUARTER, RAW, RELEASES, SERIES_SHEETS, TSJ_SHEETS,
                    WORK, canon, read_district_file, read_phase_xls, read_population_table, read_series,
                    province_key, verify_locked, write_csv, write_json, norm)

NOTE_ALIASES = {  # spellings in the CGPJ notes -> canonical district key
    "guimar": "guimar",
    "vinaros": "vinaros",
}


def build_districts():
    pop = read_population_table()
    d1 = read_district_file()
    phases = read_phase_xls()
    by_key = {canon(p["district"]): p for p in pop}
    d1_by_key = {canon(n): n for n in d1}
    unmatched_d1 = sorted(n for n in d1 if canon(n) not in by_key)
    unmatched_pop = sorted(p["district"] for p in pop if canon(p["district"]) not in d1_by_key)

    phase_of, seat_problems = {}, []
    for ph in (1, 2):
        for prov, mun in sorted(phases[ph]):
            k = canon(mun)
            if k not in by_key:
                seat_problems.append({"phase": ph, "province": prov, "seat": mun, "problem": "not a district"})
                continue
            if province_key(by_key[k]["province"]) != province_key(prov):
                seat_problems.append({"phase": ph, "province": prov, "seat": mun, "problem": "province differs"})
            if k in phase_of and phase_of[k] != ph:
                seat_problems.append({"phase": ph, "province": prov, "seat": mun, "problem": "two phases"})
            phase_of[k] = ph
    mju3 = set()
    for prov, mun in phases[3]:
        k = canon(mun)
        mju3.add(k)
        if k in phase_of:
            seat_problems.append({"phase": 3, "province": prov, "seat": mun, "problem": "MJU phase-3 seat already in phase 1/2"})
    for k in by_key:
        phase_of.setdefault(k, 3)

    rows = []
    for k, p in by_key.items():
        name = d1_by_key.get(k)
        prov = province_key(p["province"])
        rec = d1.get(name, {}) if name else {}
        tot = {y: (rec.get(y, {}).get("tot")) for y in range(2013, 2026)}
        rows.append({
            "district": name or p["district"], "district_key": k, "province": prov,
            "series_province": CEUTA_MELILLA_HOST.get(prov, prov), "ccaa": p["ccaa"],
            "phase": phase_of[k], "phase_quarter": PHASE_QUARTER[phase_of[k]],
            "in_mju_phase3_sheet": int(k in mju3), "population_2025": p["population_2025"],
            **{f"tot_{y}": tot[y] for y in range(2013, 2026)},
            **{f"{t}_{y}": rec.get(y, {}).get(t) for t in ("hip", "lau", "otr") for y in (2023, 2024, 2025)},
        })
    return rows, d1, {"unmatched_d1": unmatched_d1, "unmatched_population_table": unmatched_pop,
                      "seat_problems": seat_problems,
                      "phase_counts": {ph: sum(1 for r in rows if r["phase"] == ph) for ph in (1, 2, 3)},
                      "mju_phase3_seats": len(mju3),
                      "phase_xls_sheets": phases["sheet_names"]}


def dq1_outliers(rows):
    """DQ1 outlier rule on district-years (PROTOCOL.md §9)."""
    flags = []
    for r in rows:
        for y in range(2014, 2026):
            v = r[f"tot_{y}"]
            prev = [r[f"tot_{z}"] for z in range(2013, y) if r[f"tot_{z}"] is not None]
            if v is None or not prev:
                continue
            m = max(prev)
            jump = v >= 10 * max(1.0, m) and v - m >= 100
            drop = v == 0 and m >= 50
            if jump or drop:
                flags.append({"district": r["district"], "province": r["province"], "year": y,
                              "value": v, "max_earlier": m, "rule": "jump" if jump else "drop_to_zero"})
    return flags


def read_notes(rows):
    """The CGPJ's 'Nota falta datos' sheets, literal, matched to districts."""
    import openpyxl
    keys = {r["district_key"]: r for r in rows}
    out = []
    for quarter, fname in RELEASES:
        wb = openpyxl.load_workbook(RAW / fname, read_only=True, data_only=True)
        for sn in wb.sheetnames:
            if not re.search(r"nota falta", sn, re.I):
                continue
            title = None
            for r in wb[sn].iter_rows(values_only=True):
                v = next((x for x in r if x not in (None, "")), None)
                if v is None:
                    continue
                v = str(v).strip()
                if v.lower().startswith("partidos judiciales con falta"):
                    title = v
                    continue
                k = canon(v)
                k = NOTE_ALIASES.get(k, k)
                if k not in keys:  # try loose match: first token sequence
                    cand = [kk for kk in keys if kk.replace(" ", "") == k.replace(" ", "")
                            or kk.split(" ")[0] == k.split(" ")[0] and len(k.split(" ")[0]) > 4]
                    k = cand[0] if len(cand) == 1 else k
                d = keys.get(k)
                out.append({"release": quarter, "sheet": sn, "note_title": title, "name_in_note": v,
                            "district": d["district"] if d else "", "province": d["province"] if d else "",
                            "series_province": d["series_province"] if d else "",
                            "phase": d["phase"] if d else "",
                            "kind": ("estimated" if quarter == "2025Q3" else
                                     "partial" if quarter == "2025Q4" else "no_data")})
    return out


def weights(rows, years, flagged):
    """w_d: district share of its province's practised evictions over `years` (D1).

    A district-year flagged by the DQ1 rule is left out: the district's base is the mean of its
    unflagged years times len(years), so that it stays comparable with the other districts.
    """
    base = {}
    for r in rows:
        good = [r[f"tot_{y}"] or 0.0 for y in years if (r["district"], y) not in flagged]
        base[r["district_key"]] = (sum(good) * len(years) / len(good)) if good else 0.0
    tot = defaultdict(float)
    for r in rows:
        tot[r["series_province"]] += base[r["district_key"]]
    return {r["district_key"]: (base[r["district_key"]] / tot[r["series_province"]]
                                if tot[r["series_province"]] else 0.0) for r in rows}


def main():
    verify_locked()
    rows, d1, match = build_districts()
    flags = dq1_outliers(rows)
    flagged = {(f["district"], f["year"]) for f in flags}
    for r in rows:
        r["dq1_flags"] = ";".join(f"{f['year']}:{f['rule']}" for f in flags if f["district"] == r["district"])
    w = weights(rows, (2023, 2024), flagged)
    w24 = weights(rows, (2024,), flagged)
    for r in rows:
        r["w"] = w[r["district_key"]]
        r["w_2024"] = w24[r["district_key"]]

    notes = read_notes(rows)
    for r in rows:
        for q in ("2025Q3", "2025Q4", "2026Q1"):
            r[f"cgpj_flag_{q}"] = int(any(n["district"] == r["district"] and n["release"] == q for n in notes))

    # Province and TSJ series (latest vintage) -------------------------------------------
    prov_path = RAW / "series_provincias_1T2026_revisado.xlsx"
    tsj_path = RAW / "series_tsj_1T2026_revisado.xlsx"
    series = {sid: read_series(prov_path, sh) for sid, sh in SERIES_SHEETS.items()}
    tsj = {sid: read_series(tsj_path, sh) for sid, sh in TSJ_SHEETS.items()}
    labels = [n for n in series["P-TOT"] if n.upper() != "TOTAL"]
    label_of = {province_key(n): n for n in labels}

    B = {p: sum(series["P-TOT"][label_of[p]][f"2024Q{k}"] for k in range(1, 5)) for p in label_of}
    exposure = []
    for p in sorted(label_of):
        ds = [r for r in rows if r["series_province"] == p]
        rec = {"province": p, "label": label_of[p], "B_2024": B[p], "n_districts": len(ds)}
        for ph in (1, 2, 3):
            rec[f"n_phase{ph}"] = sum(1 for r in ds if r["phase"] == ph)
            rec[f"W{ph}"] = sum(r["w"] for r in ds if r["phase"] == ph)
            rec[f"W{ph}_2024"] = sum(r["w_2024"] for r in ds if r["phase"] == ph)
        exposure.append(rec)

    pq = []
    for sid, data in series.items():
        for name, vals in data.items():
            for q, v in vals.items():
                if q >= "2021Q1":
                    pq.append({"series": sid, "province": "TOTAL" if name.upper() == "TOTAL" else province_key(name),
                               "quarter": q, "value": v})

    out_rows = []
    for r in sorted(rows, key=lambda r: (r["province"], r["district"])):
        out_rows.append({k: r[k] for k in (
            "district", "province", "series_province", "ccaa", "phase", "phase_quarter", "in_mju_phase3_sheet",
            "tot_2023", "tot_2024", "tot_2025", "hip_2024", "hip_2025", "lau_2024", "lau_2025", "otr_2024", "otr_2025",
            "w", "w_2024", "dq1_flags", "cgpj_flag_2025Q3", "cgpj_flag_2025Q4", "cgpj_flag_2026Q1")})
    write_csv(DATA / "districts.csv", out_rows)
    write_csv(DATA / "province_exposure.csv", exposure)
    write_csv(DATA / "province_quarter.csv", pq, ["series", "province", "quarter", "value"])
    write_csv(DATA / "cgpj_notes.csv", notes)
    write_json(WORK / "panel.json", {
        "province_series": {sid: {("TOTAL" if n.upper() == "TOTAL" else province_key(n)): v for n, v in d.items()}
                            for sid, d in series.items()},
        "tsj_series": tsj,
        "districts": rows,
        "matching": match,
        "dq1_outliers": flags,
    })
    print(f"districts {len(rows)}; phases {match['phase_counts']}; unmatched D1 {match['unmatched_d1']}; "
          f"seat problems {len(match['seat_problems'])}; DQ1 flags {len(flags)}; notes {len(notes)} "
          f"(unmatched {sum(1 for n in notes if not n['district'])})")


if __name__ == "__main__":
    main()
