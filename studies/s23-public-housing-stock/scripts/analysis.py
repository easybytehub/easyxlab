#!/usr/bin/env python3
"""S23: reconciliation of the official figures for Spain's public/social housing stock and for
the EU average they are compared with.

Reads data/figures.csv, data/eu_averages.csv, data/ovs_*.csv and raw files (MIVAU stock
estimate, OECD xlsx, Eurostat JSON, INE ECV) and writes:
  data/denominators.csv     dwelling and household counts used as denominators
  data/shares_recomputed.csv each stock count divided by each denominator
  data/oecd_average_check.csv how the OECD 'EU' and 'OECD' averages are computed
  data/reconciliation.csv   each difference between figures, its explanation and category
  data/summary.json         headline numbers (checked by check_headlines.py)
"""
import csv, json, os, re, statistics, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from s23lib import RAW, DATA

OUT = {}


def rows(name):
    return list(csv.DictReader(open(os.path.join(DATA, name), encoding="utf-8")))


def write(name, rr):
    with open(os.path.join(DATA, name), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rr[0].keys()))
        w.writeheader(); w.writerows(rr)


def pct(a, b, nd=2):
    return round(100 * a / b, nd)


# ---------------------------------------------------------------- denominators
def mivau_stock():
    tot = {}
    with open(os.path.join(RAW, "mivau", "VDP002_01.csv"), encoding="utf-8-sig") as f:
        for r in csv.DictReader(f, delimiter=";"):
            k = (int(r["Año"]), r["Tipo_Vivienda"])
            tot[k] = tot.get(k, 0) + int(float(r["Valor"].replace(",", ".")))
    return tot


def ecv():
    out = {}
    with open(os.path.join(RAW, "ine", "ecv_76845.csv"), encoding="utf-8-sig") as f:
        for r in csv.DictReader(f, delimiter=";"):
            if r["Comunidades y Ciudades Autónomas"] == "Total Nacional" and r["Total"].strip():
                out[(r["Régimen de tenencia de la vivienda principal"], int(r["Periodo"]))] = float(r["Total"].replace(",", "."))
    return out


def eurostat():
    d = json.load(open(os.path.join(RAW, "eurostat", "ilc_lvho02_ES_EU27.json"), encoding="utf-8"))
    dims, sizes = d["id"], d["size"]
    idx = {k: d["dimension"][k]["category"]["index"] for k in dims}

    def get(**kw):
        pos = 0
        for k, s in zip(dims, sizes):
            pos = pos * s + idx[k][kw[k]]
        return d["value"].get(str(pos))
    out = {}
    for geo in ("ES", "EU27_2020"):
        for t in ("RENT_FR", "RENT_MKT"):
            for y in range(2018, 2026):
                out[(geo, t, y)] = get(freq="A", rskpovth="TOTAL", hhcomp="TOTAL", tenure=t, unit="PC", geo=geo, time=str(y))
    return out, d.get("updated", "")


def main():
    st = mivau_stock()
    den = []
    for y in (2019, 2021, 2023, 2024, 2025):
        den.append({"denominator": f"MIVAU stock estimate, all dwellings, {y}", "year": y, "value": st[(y, "Viviendas_Totales")],
                    "source": "cdn.mivau.gob.es VDP002_01.csv (sum of 51 provinces)"})
        den.append({"denominator": f"MIVAU stock estimate, principal dwellings, {y}", "year": y, "value": st[(y, "Vivienda_Principal")],
                    "source": "cdn.mivau.gob.es VDP002_01.csv (sum of 51 provinces)"})
    den.append({"denominator": "Households, Census 2021 (as used by OVS 2024 and Housing Europe 2025)", "year": 2021, "value": 18539223,
                "source": "Housing Europe 2025 Spain profile, total of the tenure table"})
    den.append({"denominator": "Households, as stated in Ley 12/2023 and OVS 2020", "year": "not stated", "value": 18600000, "source": "«18,6 millones de hogares»"})
    den.append({"denominator": "All dwellings implied by OECD PH4.2 for Spain (290,000 / 1.1308%)", "year": 2019, "value": round(290000 / 0.011308203126523195),
                "source": "OECD xlsx"})
    den.append({"denominator": "All dwellings, Housing Europe 2021 Spain profile («2019 data»)", "year": 2019, "value": 25793323, "source": "Housing Europe 2021"})
    write("denominators.csv", den)

    # each stock count against each denominator
    counts = [("290,000 (OVS 2019 survey)", 290000), ("318,000 (OVS 2023 survey)", 318000), ("300,000 (Banco de España, rounded)", 300000)]
    sh = []
    for lab, n in counts:
        for d in den:
            sh.append({"stock": lab, "denominator": d["denominator"], "share_pct": pct(n, d["value"])})
    write("shares_recomputed.csv", sh)
    S = {d["denominator"]: d["value"] for d in den}
    OUT["share_318k_total_2024"] = pct(318000, st[(2024, "Viviendas_Totales")], 1)
    OUT["share_318k_principal_2024"] = pct(318000, st[(2024, "Vivienda_Principal")], 1)
    OUT["share_318k_census_households"] = pct(318000, 18539223, 2)
    OUT["share_290k_households_186"] = pct(290000, 18600000, 1)
    OUT["share_290k_total_2019_current_vintage"] = pct(290000, st[(2019, "Viviendas_Totales")], 1)
    OUT["oecd_implied_total_dwellings"] = round(290000 / 0.011308203126523195)
    OUT["stock_total_2024"] = st[(2024, "Viviendas_Totales")]
    OUT["stock_principal_2024"] = st[(2024, "Vivienda_Principal")]
    OUT["share_318k_max_households"] = round(max(pct(318000, st[(y, "Vivienda_Principal")], 3) for y in (2019, 2021, 2023, 2024, 2025)), 2)
    OUT["share_318k_min_households"] = round(min(pct(318000, st[(y, "Vivienda_Principal")], 3) for y in (2019, 2021, 2023, 2024, 2025)), 2)
    OUT["ratio_total_to_principal_2024"] = round(st[(2024, "Viviendas_Totales")] / st[(2024, "Vivienda_Principal")], 2)

    # ------------------------------------------------------------ OECD averages
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(RAW, "oecd", "PH4-2-Social-rental-housing-stock.xlsx"), data_only=True)
    A = {}
    for row in wb["Table PH4.2.A1"].iter_rows(min_row=6, values_only=True):
        if not row[0]:
            continue
        name = re.sub(r"\s*\(.*", "", str(row[0])).strip()
        if isinstance(row[5], (int, float)):
            A[name] = {"n": row[4], "pct": row[5], "year": row[6], "recent": True}
        elif isinstance(row[2], (int, float)):
            A[name] = {"n": row[1] if isinstance(row[1], (int, float)) else None, "pct": row[2], "year": row[3], "recent": False}
    # the countries plotted in Figure PH4.2.1 and their plotted values. One row carries the
    # sheet's notes in its label cell; its value is Belgium's (matched to Table A1).
    fig, agg = {}, {}
    for row in wb["Figure PH4.2.1"].iter_rows(values_only=True):
        v = [x for x in row if x is not None]
        nums = [x for x in v if isinstance(x, (int, float))]
        if not nums or not isinstance(v[0], str):
            continue
        label = v[0].strip()
        if label in ("EU", "OECD"):
            agg[label] = nums[0]
            continue
        if label.startswith("Notes"):
            label = [k for k in A if abs(A[k]["pct"] - nums[0]) < 1e-9][0]
        fig[{"Slovak republic": "Slovak Republic"}.get(label, label)] = nums[0]
    EU = {"Netherlands", "Austria", "Denmark", "France", "Ireland", "Finland", "Poland", "Slovenia", "Belgium", "Czechia",
          "Hungary", "Germany", "Slovak Republic", "Italy", "Latvia", "Spain", "Estonia", "Portugal", "Lithuania", "Malta", "Luxembourg"}
    eu_plot = sorted(k for k in fig if k in EU)
    all_plot = sorted(fig)
    eu_mean = statistics.mean(fig[k] for k in eu_plot)
    oecd_mean = statistics.mean(fig[k] for k in all_plot)
    eu_w = sum(A[k]["n"] for k in eu_plot) / sum(A[k]["n"] / (A[k]["pct"] / 100) for k in eu_plot) * 100
    w_set = [k for k in all_plot if A.get(k, {}).get("n")]
    oecd_w = sum(A[k]["n"] for k in w_set) / sum(A[k]["n"] / (A[k]["pct"] / 100) for k in w_set) * 100
    all_recent_n = sum(A[k]["n"] for k in A if A[k]["recent"])
    chk = [
        {"average": "EU (OECD printed)", "value": round(agg["EU"], 6), "n_countries": "", "countries": ""},
        {"average": "EU, unweighted mean of the EU members plotted in Figure PH4.2.1", "value": round(eu_mean, 6), "n_countries": len(eu_plot), "countries": "; ".join(eu_plot)},
        {"average": "EU, weighted (social dwellings / all dwellings, Table A1), same countries", "value": round(eu_w, 6), "n_countries": len(eu_plot), "countries": ""},
        {"average": "OECD (OECD printed)", "value": round(agg["OECD"], 6), "n_countries": "", "countries": ""},
        {"average": "OECD, unweighted mean of all countries plotted in Figure PH4.2.1 (Colombia included)", "value": round(oecd_mean, 6), "n_countries": len(all_plot), "countries": "; ".join(all_plot)},
        {"average": "OECD, weighted, same countries", "value": round(oecd_w, 6), "n_countries": len(w_set), "countries": ""},
        {"average": "Social rental dwellings, sum of all countries with data around 2022 (Table A1)", "value": round(all_recent_n), "n_countries": sum(1 for k in A if A[k]["recent"]), "countries": ""},
    ]
    write("oecd_average_check.csv", chk)
    OUT["oecd_eu_printed"] = round(agg["EU"], 2)
    OUT["oecd_oecd_printed"] = round(agg["OECD"], 2)
    OUT["oecd_eu_unweighted"] = round(eu_mean, 2)
    OUT["oecd_eu_exact"] = abs(eu_mean - agg["EU"]) < 1e-9
    OUT["oecd_oecd_exact"] = abs(oecd_mean - agg["OECD"]) < 1e-9
    OUT["oecd_eu_n"] = len(eu_plot)
    OUT["oecd_eu_weighted"] = round(eu_w, 1)
    OUT["oecd_oecd_unweighted"] = round(oecd_mean, 2)
    OUT["oecd_oecd_weighted"] = round(oecd_w, 1)
    OUT["oecd_oecd_n"] = len(all_plot)
    OUT["oecd_sum_dwellings_million"] = round(all_recent_n / 1e6, 1)
    OUT["oecd_spain_pct"] = round(A["Spain"]["pct"], 2)
    OUT["oecd_spain_year"] = A["Spain"]["year"]

    # ------------------------------------------------------------ OVS tables
    eut = rows("ovs_eu_table.csv")
    u28 = [r for r in eut if r["bulletin"] == "2020" and r["country"].startswith("UE")][0]
    u27 = [r for r in eut if r["bulletin"] == "2024" and r["country"].startswith("UE")][0]
    uk = [r for r in eut if r["bulletin"] == "2020" and r["country"] == "Reino Unido"][0]
    OUT["ovs2020_eu28_pct"] = round(int(u28["social_dwellings"]) / int(u28["principal_dwellings"]) * 100, 1)
    OUT["ovs2024_eu27_pct"] = round(int(u27["social_dwellings"]) / int(u27["principal_dwellings"]) * 100, 1)
    OUT["ovs2020_eu_without_uk_pct"] = round((int(u28["social_dwellings"]) - int(uk["social_dwellings"])) /
                                             (int(u28["principal_dwellings"]) - int(uk["principal_dwellings"])) * 100, 1)
    OUT["ovs2020_uk_social_dwellings"] = int(uk["social_dwellings"])
    es20 = [r for r in eut if r["bulletin"] == "2020" and r["country"] == "España"][0]
    es24 = [r for r in eut if r["bulletin"] == "2024" and r["country"] == "España"][0]
    OUT["ovs2020_spain_eu_table_pct"] = float(es20["social_share_of_principal_pct"])
    OUT["ovs2024_spain_eu_table_pct"] = float(es24["social_share_of_principal_pct"])
    n17 = sum(1 for r in eut if r["bulletin"] == "2024" and not r["country"].startswith("UE"))
    OUT["ovs2024_eu_table_countries"] = n17

    reg = rows("ovs_regions.csv")
    r19 = [r for r in reg if r["source"].startswith("OVS 2020") and r["region"] == "TOTAL"][0]
    r23 = [r for r in reg if r["source"].startswith("OVS 2024") and r["region"] == "TOTAL"][0]
    OUT["regional_rental_2019"] = int(r19["rental"])
    OUT["regional_rental_2023"] = int(r23["rental"])
    OUT["regional_rental_ppp_2023"] = int(r23["rental_ppp"])
    OUT["regional_rental_growth_pct"] = round((int(r23["rental"]) / int(r19["rental"]) - 1) * 100, 1)
    OUT["regional_rental_public_2023"] = int(r23["rental_public"])
    OUT["regional_rental_public_growth_pct"] = round((int(r23["rental_public"]) / int(r19["rental"]) - 1) * 100, 1)
    t23 = [r for r in rows("ovs_regions_t23.csv") if r["region"] == "TOTAL"][0]
    OUT["t23_printed_rental_2023"] = int(t23["rental_2023"])
    OUT["t23_printed_growth_pct"] = round((int(t23["rental_2023"]) / int(t23["rental_2019"]) - 1) * 100, 1)
    body = [r for r in reg if r["source"].startswith("OVS 2024") and r["region"] != "TOTAL"]
    by = {r["region"]: int(r["rental"]) for r in body}
    diff = int(r23["rental"]) - int(t23["rental_2023"])
    # which rows does the printed total leave out?
    import itertools
    omitted = None
    for k in (1, 2, 3):
        for combo in itertools.combinations(sorted(by), k):
            if sum(by[c] for c in combo) == diff:
                omitted = combo; break
        if omitted:
            break
    OUT["t23_rental_gap"] = diff
    OUT["t23_rental_gap_rows"] = list(omitted) if omitted else []
    tot_body = {r["region"]: int(r["total"]) for r in body}
    OUT["t22_total_2023"] = int(r23["total"])
    OUT["text_total_2023"] = 250372
    OUT["text_gap_equals_ceuta_total"] = (int(r23["total"]) - 250372) == tot_body["Ceuta"]

    mun = rows("ovs_municipal_2023.csv")
    d = [m for m in mun if m["has_data"]]
    OUT["mun_listed"] = len(mun)
    OUT["mun_with_data"] = len(d)
    OUT["mun_with_data_2023_survey"] = sum(1 for m in d if not m["data_from_2019_survey"])
    OUT["mun_rental_observed"] = sum(int(m["rental"] or 0) for m in d)
    OUT["mun_population_observed"] = sum(int(m["population_2023"]) for m in d)
    pop_es = int([r for r in eut if r["bulletin"] == "2024" and r["country"] == "España"][0]["population"])
    OUT["population_spain_ovs2024"] = pop_es
    OUT["mun_scaled_by_population"] = round(OUT["mun_rental_observed"] * pop_es / OUT["mun_population_observed"])
    OUT["mun_observed_share_of_121k_pct"] = round(OUT["mun_rental_observed"] / 121000 * 100)
    OUT["mun_with_data_2019_carryover"] = OUT["mun_with_data"] - OUT["mun_with_data_2023_survey"]
    OUT["mun_scaled_bulletin_base"] = round(OUT["mun_rental_observed"] * pop_es / 24500000)   # the bulletin's «24,5 millones»
    cm = [m for m in mun if m["municipality"] in ("Ceuta", "Melilla")]
    OUT["ceuta_melilla_in_municipal_table"] = sum(int(m["rental"] or 0) for m in cm)
    OUT["ceuta_melilla_in_regional_table"] = by.get("Ceuta", 0) + by.get("Melilla", 0)
    bcn = [m for m in mun if m["municipality"] == "Barcelona"][0]
    OUT["barcelona_table"] = int(bcn["rental"])
    OUT["barcelona_text"] = 8758
    OUT["ceuta_melilla_scaled"] = round(OUT["ceuta_melilla_in_municipal_table"] * pop_es / OUT["mun_population_observed"])
    OUT["ceuta_melilla_scaled_share_pct"] = round(OUT["ceuta_melilla_scaled"] / 318000 * 100, 1)

    # ------------------------------------------------------------ surveys
    e = ecv()
    OUT["ecv_below_market"] = {y: e[("Alquiler inferior al precio de mercado", y)] for y in range(2012, 2026)}
    OUT["ecv_cesion_2025"] = e[("Cesión", 2025)]
    OUT["ecv_25_years"] = [y for y in range(2004, 2026) if e.get(("Alquiler inferior al precio de mercado", y)) == 2.5]
    OUT["ecv_34_years"] = [y for y in range(2004, 2026) if e.get(("Alquiler inferior al precio de mercado", y)) == 3.4]
    OUT["ecv_below_plus_cesion_2025"] = round(e[("Alquiler inferior al precio de mercado", 2025)] + e[("Cesión", 2025)], 1)
    es, upd = eurostat()
    OUT["silc_es_rent_fr"] = {y: es[("ES", "RENT_FR", y)] for y in (2019, 2023, 2025)}
    OUT["silc_eu_rent_fr"] = {y: es[("EU27_2020", "RENT_FR", y)] for y in (2018, 2019, 2023, 2025)}
    OUT["eurostat_updated"] = upd

    # ------------------------------------------------------------ figures counts
    fig_es, fig_eu = rows("figures.csv"), rows("eu_averages.csv")
    OUT["n_spain_rows"] = len(fig_es)
    OUT["n_eu_rows"] = len(fig_eu)
    OUT["n_documents"] = len({r["raw"] for r in fig_es + fig_eu})   # one raw source per document
    OUT["n_quotes_verified"] = sum(1 for r in fig_es + fig_eu if r["quote_verified"] == "yes")

    # ------------------------------------------------------------ reconciliation table
    # status rule: "explained" only if every number in the row reproduces from published data;
    # "partly explained" if the concept is explained but a number does not reproduce exactly.
    R = []

    def rec(ids, figures, explanation, category, status, check=""):
        assert status in ("explained", "partly explained", "not explained", "error in a source")
        R.append({"figures": ids, "values": figures, "explanation": explanation,
                  "category": category, "status": status, "check": check})
    ecvb = OUT["ecv_below_market"]
    rec("ES-01/02, ES-03/04, ES-14", "290,000 = 1.6% of households; = 1.1% of dwellings",
        "Denominator. 1.6% is a share of 18.6 million households; Housing Europe 2021 divides the same 290,000 by 25,793,323 dwellings.",
        "denominator", "explained", f"290,000/18.6M = {OUT['share_290k_households_186']}%; 290,000/25,793,323 = {pct(290000, 25793323)}%")
    rec("ES-16, ES-17, ES-18", "OECD 290,000 = 1.13% of dwellings (2019); Banco de España ~300,000 = 1.5% of principal dwellings, 1% of all dwellings",
        f"Denominator and rounding. The OECD's 1.13% implies {OUT['oecd_implied_total_dwellings']:,} dwellings, a stock figure we did not find in any source (Housing Europe 2021: 25,793,323; the Ministry's current series for 2019: {S['MIVAU stock estimate, all dwellings, 2019']:,}). The Banco de España's rounded figures are close to 300,000 over 2023 principal dwellings ({pct(300000, S['MIVAU stock estimate, principal dwellings, 2023'])}%) and all dwellings ({pct(300000, S['MIVAU stock estimate, all dwellings, 2023'])}%).",
        "denominator", "partly explained", "")
    rec("ES-07, ES-22, ES-24", "318,000 = 1.72% of households; '1.5-1.7% of households'; '<2%' of the housing supply",
        f"Denominator. 318,000 / 18.5 million Census households = {OUT['share_318k_census_households']}%; / 27.0 million dwellings (2024) = {OUT['share_318k_total_2024']}%. The Commission's lower bound of 1.5% is not reproduced: over the Ministry's principal dwellings for 2019-2025, 318,000 gives {OUT['share_318k_min_households']}-{OUT['share_318k_max_households']}%. It equals the Banco de España's share for ~300,000.",
        "denominator", "partly explained", "")
    rec("ES-03/04 vs ES-06/07; ES-01/02 vs ES-09", "290,000 (2019 survey) vs 318,000 (2023 survey)",
        f"Reference year and survey round. Part of the change is scope: the 2023 regional count adds a column of {OUT['regional_rental_ppp_2023']:,} dwellings under public-private schemes that the 2019 table does not show, and most municipal figures in 2023 are carried over from 2019. The bulletin does not decompose the change.",
        "reference year / survey round", "partly explained", f"regional rental {OUT['regional_rental_2019']:,} (2019) vs {OUT['regional_rental_2023']:,} (2023), of which {OUT['regional_rental_ppp_2023']:,} public-private")
    rec("ES-09/10, ES-08, ES-11, ES-20", "1.7% vs 3.3% / 3.5%",
        "Definition. 1.7% counts dwellings owned by regions and municipalities. 3.3% (2023) and 3.5% (2025) are INE's ECV shares of households paying a rent below the market price, any landlord; free accommodation excluded.",
        "definition", "explained", f"ECV below-market rent 2023 = {ecvb[2023]}, 2025 = {ecvb[2025]}")
    rec("ES-05, ES-13", "Spain 2.5% (OVS 2020 EU table; Housing Europe 2019)",
        f"Definition and year. Neither document names its source; 2.5 equals INE's ECV below-market-rent share in {OUT['ecv_25_years'][0]}-{OUT['ecv_25_years'][-1]} (every year), and the 2024 edition of the OVS table states that Spain's cell is the ECV share.",
        "definition", "explained", f"ECV below-market rent = 2.5 in {OUT['ecv_25_years']}")
    rec("ES-23, ES-15", "2.5-3.4% (Commission 2026); 3.4% (Housing Europe 2025)",
        f"Definition and year. Both bounds equal ECV below-market-rent shares (2.5 in 2012-2016; 3.4 in {OUT['ecv_34_years']}). The Commission describes the range as 'protected or regulated housing (rental and sale)' and Housing Europe as protected housing; the ECV measures households paying a below-market rent.",
        "definition", "explained", "")
    rec("ES-11 vs ES-12", "3.5% (households, below-market rent) vs 9.0% (persons, reduced price or free)",
        f"Definition and unit. Eurostat's category adds free accommodation (INE 'Cesión', {OUT['ecv_cesion_2025']}% of households in 2025) and counts persons. Below-market rent plus cession is {OUT['ecv_below_plus_cesion_2025']}% of households, not 9.0% of persons; the unit difference is not reproduced.",
        "definition", "partly explained", f"EU-SILC ES RENT_FR {OUT['silc_es_rent_fr']}")
    rec("ES-19, ES-21, EU-06, EU-08", "Commission 2025: «1.5% the total stock ... EU average of 9%»; Council 2025: «1,5 % of the total housing stock ... EU average of 9 %», both attributed to the Banco de España",
        "Attribution. The Commission cites a Banco de España paper (footnote 4); the Council writes 'the Bank of Spain' without a reference. The paper gives 1.5% of principal dwellings, 1% of all dwellings and an average of about 7%; we found no Banco de España text with 9%. No other origin is stated for these texts' 9%.",
        "attribution", "error in a source", "")
    rec("EU-01 vs EU-02", "9.3% (OVS 2020) vs 8.0% (OVS 2024)",
        f"Scope: EU-28 with the United Kingdom ({OUT['ovs2020_uk_social_dwellings']:,} social dwellings) vs EU-27. Removing the UK from the 2020 table gives {OUT['ovs2020_eu_without_uk_pct']}%.",
        "scope", "explained", f"recomputed: EU-28 {OUT['ovs2020_eu28_pct']}%, EU-27 {OUT['ovs2024_eu27_pct']}%, EU-28 without UK {OUT['ovs2020_eu_without_uk_pct']}%")
    rec("EU-03 (OECD 7% / EU 8%)", "7% vs 8%",
        f"Country set: unweighted means of the countries plotted in Figure PH4.2.1, all {OUT['oecd_oecd_n']} with Colombia ({OUT['oecd_oecd_unweighted']}) and the {OUT['oecd_eu_n']} EU members ({OUT['oecd_eu_unweighted']}). Both reproduce exactly.",
        "method", "explained", f"printed EU {OUT['oecd_eu_printed']}, OECD {OUT['oecd_oecd_printed']}")
    rec("EU-04, EU-05, EU-09, EU-10, EU-11", "about 7%; 6-7%; 'around 28 million homes'",
        f"'About 7%' matches the OECD mean. '6-7%' of the EU stock is close to the 19 EU countries weighted by dwellings ({OUT['oecd_eu_weighted']}%), but the texts do not say how they averaged. SWD(2025) 1053 pairs 6-7% with 'around 28 million homes', the total for all countries with recent OECD data ({OUT['oecd_sum_dwellings_million']} million recomputed).",
        "method", "partly explained", "")
    rec("EU-07 (RD 326/2026)", "EU average 9% in the 2026-2030 Plan",
        "No source stated. The 9% values we could trace are the OVS 2020 EU-28 average (9.3%) and Eurostat's EU-27 share of persons renting at a reduced price or free in 2018-2019 (9.3%, 9.1%); the Observatory's 2024 bulletin, dated January 2025, gives 8.0% for the EU-27. We cannot tell which, if any, the Plan used.",
        "unexplained", "not explained", "")
    rec("ES-08 + EU-02 (OVS 2024 Tabla 2.1)", "Spain 3.3% vs EU-27 8.0%",
        "Definition: Spain's cell is the ECV below-market-rent share of households; the other countries' cells are Housing Europe social-housing shares, weighted by 2011 principal dwellings. Both numbers reproduce; the comparison is not like-for-like.",
        "definition", "explained", "")
    rec("OVS 2024, Tabla 2.2 vs Tabla 2.3 and text", f"regional rental stock 2023: {OUT['regional_rental_2023']:,} (sum of rows) vs {OUT['t23_printed_rental_2023']:,} (printed total); all tenures 251,937 vs 250,372 (text)",
        f"Arithmetic: the printed rental total equals the sum of the regions without {' and '.join(OUT['t23_rental_gap_rows'])}; the text's 250,372 equals the table total without Ceuta; the for-sale and 'other' totals of Tabla 2.3 also omit Ceuta. The text's 'aumento del 5%' is consistent with the printed total ({OUT['t23_printed_growth_pct']}%). Over all regions the rental columns give {OUT['regional_rental_growth_pct']}%, of which {OUT['regional_rental_ppp_2023']:,} dwellings are in a public-private column new in 2023; the public-ownership column alone gives {OUT['regional_rental_public_growth_pct']}%.",
        "arithmetic", "error in a source", "")
    rec("OVS 2024, regional and municipal tables", f"Ceuta and Melilla: {OUT['ceuta_melilla_in_regional_table']:,} rental dwellings in both tables",
        f"Possible double count. The two autonomous cities appear with the same figures in the regional table (part of 197,000) and the municipal table (the sample scaled to 121,000). If both entered the total, the second count adds about {OUT['ceuta_melilla_scaled']:,} after scaling ({OUT['ceuta_melilla_scaled_share_pct']}% of 318,000). The bulletin does not publish the scaling.",
        "unexplained", "not explained", "")
    rec("OVS 2024 municipal component", f"121,000 vs {OUT['mun_rental_observed']:,} observed",
        f"Method: {OUT['mun_with_data']} municipalities of over 20,000 inhabitants have figures in Tabla 2.8 ({OUT['mun_with_data_2023_survey']} from the 2023 survey, {OUT['mun_with_data_2019_carryover']} carried over from 2019), {OUT['mun_population_observed']/1e6:.1f} million people. Scaling their {OUT['mun_rental_observed']:,} rental dwellings to Spain's population gives {OUT['mun_scaled_by_population']:,}; with the bulletin's own base (168 municipalities, 24.5 million) it gives {OUT['mun_scaled_bulletin_base']:,}. Our sample is not the one the bulletin describes.",
        "method", "partly explained", f"scaled {OUT['mun_scaled_by_population']:,} / {OUT['mun_scaled_bulletin_base']:,} vs printed 121,000")
    write("reconciliation.csv", R)
    OUT["n_reconciliation_rows"] = len(R)
    cats, stat = {}, {}
    for r in R:
        cats[r["category"]] = cats.get(r["category"], 0) + 1
        stat[r["status"]] = stat.get(r["status"], 0) + 1
    OUT["reconciliation_categories"] = cats
    OUT["reconciliation_status"] = stat
    json.dump(OUT, open(os.path.join(DATA, "summary.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1, sort_keys=True)
    for k, v in OUT.items():
        print(k, v)


if __name__ == "__main__":
    main()
